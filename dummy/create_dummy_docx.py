import docx
from docx.shared import Pt, RGBColor, Inches

doc = docx.Document()

# Set standard margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Header
header = doc.sections[0].header
hp = header.paragraphs[0]
hrun = hp.add_run("Astra Life Document Template - Confidential")
hrun.font.name = "Arial"
hrun.font.size = Pt(8.5)

# Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("SURAT PENGANTAR POLIS ASURANSI JIWA ASTRA")
title_run.font.name = "Arial"
title_run.font.size = Pt(16)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0, 44, 108)

# Paragraph
p1 = doc.add_paragraph()
r1 = p1.add_run("Kepada Yth. Pemegang Polis,\nTerima kasih telah mempercayakan perlindungan masa depan Anda dan keluarga kepada PT Asuransi Jiwa Astra (Astra Life). Bersama surat ini kami lampirkan dokumen ikhtisar pertanggungan asuransi jiwa Anda.")
r1.font.name = "Arial"
r1.font.size = Pt(11)

# Table
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Informasi"
hdr_cells[1].text = "Keterangan"

for cell in hdr_cells:
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"
            r.font.bold = True

data = [
    ("Nama Perusahaan", "PT Asuransi Jiwa Astra"),
    ("Jenis Polis", "Asuransi Jiwa Tradisional"),
    ("Status Font Default", "Arial (Harus dikonversi ke Roboto)"),
]

for info, desc in data:
    row_cells = table.add_row().cells
    r0 = row_cells[0].paragraphs[0].add_run(info)
    r0.font.name = "Arial"
    r1 = row_cells[1].paragraphs[0].add_run(desc)
    r1.font.name = "Arial"

doc.save("dummy/template.docx")
print("dummy/template.docx created successfully.")
