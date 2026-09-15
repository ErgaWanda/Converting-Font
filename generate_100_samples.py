"""
Script untuk menghasilkan 100 dokumen sampel (50 DOCX + 50 XML) dengan font Arial
untuk menguji fitur Batch & Folder Font Converter Astra Life.
"""

import os
import shutil
import zipfile
import docx
from docx.shared import Pt, RGBColor, Inches

OUTPUT_DIR = "dokumen_arial_100"
ZIP_OUTPUT = "dokumen_arial_100.zip"

def create_100_documents():
    if os.path.exists(OUTPUT_DIR):
        shutil.rmtree(OUTPUT_DIR)
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    names = [
        "Budi Santoso", "Siti Rahmawati", "Ahmad Fauzi", "Dewi Lestari", "Rian Hidayat",
        "Maya Anggraini", "Hendra Wijaya", "Indah Permatasari", "Eko Prasetyo", "Putri Utami",
        "Bambang Setiawan", "Nurul Hidayati", "Agus Supriyanto", "Dian Sastro", "Reza Rahadian",
        "Chandra Kirana", "Fajar Ramadhan", "Gita Gutawa", "Haris Munandar", "Ika Safitri",
        "Joko Widodo", "Kartika Sari", "Lukman Hakim", "Mega Suryani", "Nadia Vega",
        "Oki Setiana", "Pandu Pratama", "Qori Sandioriva", "Rizky Febian", "Syahrini Fatimah",
        "Taufik Hidayat", "Umar Bakri", "Vina Panduwinata", "Wahyu Hidayat", "Xaverius Tedy",
        "Yani Maryani", "Zainal Abidin", "Aditya Nugraha", "Bella Safira", "Cecep Reza",
        "Danang Pradana", "Erna Susanti", "Farhan Ali", "Gatot Subroto", "Hany Puspita",
        "Irfan Bachdim", "Julia Perez", "Kiki Amalia", "Lilis Karlina", "Muammar ZA"
    ]

    products = [
        "Astra Life AVA iPrime Protection",
        "Astra Life AVA iFuture Assurance",
        "Astra Life AVA Health Optima",
        "Astra Life AVA Critical Care",
        "Astra Life AVA Investa Maksima"
    ]

    print("Membuat 50 file DOCX berfont Arial...")
    for i in range(1, 51):
        name = names[i - 1]
        prod = products[(i - 1) % len(products)]
        no_polis = f"POLIS-DOCX-{20260000 + i}"
        premi = f"Rp {(1_500_000 + (i * 250_000)):,}".replace(",", ".")
        up = f"Rp {(100_000_000 + (i * 25_000_000)):,}".replace(",", ".")

        doc = docx.Document()


        header = doc.sections[0].header
        hp = header.paragraphs[0]
        hrun = hp.add_run(f"Astra Life Insurance - Policy Doc #{no_polis} [Strictly Confidential]")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)


        p_title = doc.add_paragraph()
        r_title = p_title.add_run(f"IKHTISAR POLIS ASURANSI - DOKUMEN #{i:03d}")
        r_title.font.name = "Arial"
        r_title.font.size = Pt(16)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0, 44, 108)


        p_intro = doc.add_paragraph()
        r_intro = p_intro.add_run(
            f"Kepada Yth. Bapak/Ibu {name},\n"
            f"Terima kasih atas kepercayaan Anda memilih {prod} sebagai solusi proteksi finansial keluarga. "
            f"Berikut ringkasan data polis resmi yang tercatat di sistem PT Asuransi Jiwa Astra:"
        )
        r_intro.font.name = "Arial"
        r_intro.font.size = Pt(11)


        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Parameter Dokumen"
        hdr_cells[1].text = "Detail Polis Nasabah"

        for cell in hdr_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.bold = True

        data_rows = [
            ("Nomor Polis", no_polis),
            ("Nama Pemegang Polis", name),
            ("Produk Asuransi", prod),
            ("Premi Tahunan", premi),
            ("Uang Pertanggungan", up),
            ("Status Polis", "AKTIF (In-Force)"),
            ("Font Standar Dokumen", "Arial (Wajib konversi ke Roboto)")
        ]

        for label, val in data_rows:
            row_cells = table.add_row().cells
            r0 = row_cells[0].paragraphs[0].add_run(label)
            r0.font.name = "Arial"
            r1 = row_cells[1].paragraphs[0].add_run(val)
            r1.font.name = "Arial"


        p_footer = doc.add_paragraph()
        r_footer = p_footer.add_run(
            "\nCatatan: Dokumen ini dibuat otomatis oleh Document Automation Engine Astra Life. "
            "Seluruh format tipografi menggunakan standar perusahaan."
        )
        r_footer.font.name = "Arial"
        r_footer.font.size = Pt(9)
        r_footer.font.italic = True

        filename = f"{OUTPUT_DIR}/Dokumen_{i:03d}_{no_polis}_{name.replace(' ', '_')}.docx"
        doc.save(filename)

    print("Membuat 50 file XML berfont Arial...")
    xml_template = """<?xml version="1.0" encoding="UTF-8"?>
<document version="2.0" xmlns="http://schemas.astralife.co.id/docgen/2026">
    <metadata>
        <title>Sertifikat Kepesertaan Asuransi - {NO_POLIS}</title>
        <author>PT Asuransi Jiwa Astra</author>
        <default-font family="Arial" size="11pt" color="#333333" />
    </metadata>
    <header>
        <company font-family="Arial" font-size="14pt" font-weight="bold" color="#002C6C">PT ASURANSI JIWA ASTRA</company>
        <sub-info font-family="Arial" font-size="9pt" color="#666666">Pondok Indah Office Tower 3, Jakarta Selatan</sub-info>
    </header>
    <body>
        <heading level="1" font-family="Arial" font-size="16pt" font-weight="bold" color="#002C6C">SERTIFIKAT KEPESERTAAN RESMI</heading>
        <paragraph font-family="Arial" font-size="10pt">
            Dokumen ini menyatakan bahwa nasabah di bawah ini sah terdaftar dalam program proteksi Astra Life:
        </paragraph>
        <policy-details font-family="Arial">
            <field name="Nomor Polis" value="{NO_POLIS}" font-family="Arial" />
            <field name="Nama Lengkap" value="{NAMA}" font-family="Arial" />
            <field name="Produk" value="{PRODUK}" font-family="Arial" />
            <field name="Premi" value="{PREMI}" font-family="Arial" />
            <field name="Uang Pertanggungan" value="{UP}" font-family="Arial" />
            <field name="Font Dokumen" value="Arial" font-family="Arial" />
        </policy-details>
        <terms font-family="Arial" font-size="8.5pt" color="#777777">
            <clause font-family="Arial">Klaim dapat diajukan secara online melalui aplikasi Astra Life.</clause>
            <clause font-family="Arial">Seluruh hak dan kewajiban mengacu pada polis induk.</clause>
        </terms>
    </body>
</document>"""

    for i in range(1, 51):
        idx = i + 50
        name = names[i - 1]
        prod = products[(i - 1) % len(products)]
        no_polis = f"POLIS-XML-{20260000 + idx}"
        premi = f"Rp {(2_000_000 + (i * 300_000)):,}".replace(",", ".")
        up = f"Rp {(200_000_000 + (i * 50_000_000)):,}".replace(",", ".")

        rendered_xml = xml_template.format(
            NO_POLIS=no_polis,
            NAMA=name,
            PRODUK=prod,
            PREMI=premi,
            UP=up
        )

        filename = f"{OUTPUT_DIR}/Dokumen_{idx:03d}_{no_polis}_{name.replace(' ', '_')}.xml"
        with open(filename, "w", encoding="utf-8") as f:
            f.write(rendered_xml)


    print(f"Mengompres folder '{OUTPUT_DIR}' menjadi '{ZIP_OUTPUT}'...")
    with zipfile.ZipFile(ZIP_OUTPUT, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(OUTPUT_DIR):
            for file in files:
                full_path = os.path.join(root, file)
                rel_path = os.path.relpath(full_path, OUTPUT_DIR)
                zf.write(full_path, rel_path)

    total_files = len(os.listdir(OUTPUT_DIR))
    print(f"\n[SUKSES] Total {total_files} file berhasil dibuat di folder '{OUTPUT_DIR}'!")
    print(f"[SUKSES] File ZIP '{ZIP_OUTPUT}' juga siap digunakan untuk pengujian upload di Web!")

if __name__ == "__main__":
    create_100_documents()
