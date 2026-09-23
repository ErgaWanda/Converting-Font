import io
import re
import os
import zipfile
import urllib.request
import xml.etree.ElementTree as ET
from typing import Optional

from fastapi import FastAPI, File, UploadFile, HTTPException, Form
from fastapi.responses import HTMLResponse, StreamingResponse, FileResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.requests import Request
import pandas as pd
import docx
from docx.oxml.ns import qn
import pymupdf
import pymupdf as fitz

app = FastAPI(
    title="Tools Converting Font",
    description="Portal Otomatisasi Dokumen: Font Converter & Batch Mapping Merger (XML, DOCX, & PDF)",
    version="1.3.0"
)


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
SAMPLES_DIR = os.path.join(BASE_DIR, "assets", "samples")
FONTS_DIR = os.path.join(BASE_DIR, "assets", "fonts")

os.makedirs(TEMPLATES_DIR, exist_ok=True)
os.makedirs(SAMPLES_DIR, exist_ok=True)
os.makedirs(FONTS_DIR, exist_ok=True)

templates = Jinja2Templates(directory=TEMPLATES_DIR)

FONTS_REGISTRY = {
    'roboto': {
        'id': 'roboto',
        'name': 'Roboto',
        'family': 'Roboto',
        'prefix': 'Roboto_',
        'badge': 'Google Font (Bawaan)',
        'description': 'Font resmi standar dengan keterbacaan tinggi dan proporsi seimbang.',
        'files': {
            'regular': os.path.join(FONTS_DIR, 'Roboto-Regular.ttf'),
            'bold': os.path.join(FONTS_DIR, 'Roboto-Bold.ttf'),
            'italic': os.path.join(FONTS_DIR, 'Roboto-Italic.ttf'),
            'bolditalic': os.path.join(FONTS_DIR, 'Roboto-BoldItalic.ttf')
        },
        'urls': {
            'Roboto-Regular.ttf': 'https://raw.githubusercontent.com/googlefonts/roboto/main/src/hinted/Roboto-Regular.ttf',
            'Roboto-Bold.ttf': 'https://raw.githubusercontent.com/googlefonts/roboto/main/src/hinted/Roboto-Bold.ttf',
            'Roboto-Italic.ttf': 'https://raw.githubusercontent.com/googlefonts/roboto/main/src/hinted/Roboto-Italic.ttf',
            'Roboto-BoldItalic.ttf': 'https://raw.githubusercontent.com/googlefonts/roboto/main/src/hinted/Roboto-BoldItalic.ttf'
        }
    },
    'opensans': {
        'id': 'opensans',
        'name': 'Open Sans',
        'family': 'Open Sans',
        'prefix': 'OpenSans_',
        'badge': 'Google Font',
        'description': 'Font modern humanis sans-serif yang bersih dan mudah dibaca.',
        'files': {
            'regular': os.path.join(FONTS_DIR, 'OpenSans-Regular.ttf'),
            'bold': os.path.join(FONTS_DIR, 'OpenSans-Bold.ttf'),
            'italic': os.path.join(FONTS_DIR, 'OpenSans-Italic.ttf'),
            'bolditalic': os.path.join(FONTS_DIR, 'OpenSans-BoldItalic.ttf')
        },
        'urls': {
            'OpenSans-Regular.ttf': 'https://raw.githubusercontent.com/googlefonts/opensans/main/fonts/ttf/OpenSans-Regular.ttf',
            'OpenSans-Bold.ttf': 'https://raw.githubusercontent.com/googlefonts/opensans/main/fonts/ttf/OpenSans-Bold.ttf',
            'OpenSans-Italic.ttf': 'https://raw.githubusercontent.com/googlefonts/opensans/main/fonts/ttf/OpenSans-Italic.ttf',
            'OpenSans-BoldItalic.ttf': 'https://raw.githubusercontent.com/googlefonts/opensans/main/fonts/ttf/OpenSans-BoldItalic.ttf'
        }
    },
    'montserrat': {
        'id': 'montserrat',
        'name': 'Montserrat',
        'family': 'Montserrat',
        'prefix': 'Montserrat_',
        'badge': 'Google Font',
        'description': 'Font geometris sans-serif kontemporer terinspirasi tipografi urban.',
        'files': {
            'regular': os.path.join(FONTS_DIR, 'Montserrat-Regular.ttf'),
            'bold': os.path.join(FONTS_DIR, 'Montserrat-Bold.ttf'),
            'italic': os.path.join(FONTS_DIR, 'Montserrat-Italic.ttf'),
            'bolditalic': os.path.join(FONTS_DIR, 'Montserrat-BoldItalic.ttf')
        },
        'urls': {
            'Montserrat-Regular.ttf': 'https://raw.githubusercontent.com/JulietaUla/Montserrat/master/fonts/ttf/Montserrat-Regular.ttf',
            'Montserrat-Bold.ttf': 'https://raw.githubusercontent.com/JulietaUla/Montserrat/master/fonts/ttf/Montserrat-Bold.ttf',
            'Montserrat-Italic.ttf': 'https://raw.githubusercontent.com/JulietaUla/Montserrat/master/fonts/ttf/Montserrat-Italic.ttf',
            'Montserrat-BoldItalic.ttf': 'https://raw.githubusercontent.com/JulietaUla/Montserrat/master/fonts/ttf/Montserrat-BoldItalic.ttf'
        }
    },
    'arial': {
        'id': 'arial',
        'name': 'Arial',
        'family': 'Arial',
        'prefix': 'Arial_',
        'badge': 'System Font',
        'description': 'Font sans-serif klasik universal yang kompatibel di semua platform.',
        'files': {
            'regular': 'C:/Windows/Fonts/arial.ttf',
            'bold': 'C:/Windows/Fonts/arialbd.ttf',
            'italic': 'C:/Windows/Fonts/ariali.ttf',
            'bolditalic': 'C:/Windows/Fonts/arialbi.ttf'
        },
        'urls': {}
    },
    'times': {
        'id': 'times',
        'name': 'Times New Roman',
        'family': 'Times New Roman',
        'prefix': 'Times_',
        'badge': 'System Font',
        'description': 'Font serif formal klasik standar untuk surat dan dokumen legal.',
        'files': {
            'regular': 'C:/Windows/Fonts/times.ttf',
            'bold': 'C:/Windows/Fonts/timesbd.ttf',
            'italic': 'C:/Windows/Fonts/timesi.ttf',
            'bolditalic': 'C:/Windows/Fonts/timesbi.ttf'
        },
        'urls': {}
    },
    'calibri': {
        'id': 'calibri',
        'name': 'Calibri',
        'family': 'Calibri',
        'prefix': 'Calibri_',
        'badge': 'System Font',
        'description': 'Font sans-serif modern Microsoft Office dengan sudut membulat.',
        'files': {
            'regular': 'C:/Windows/Fonts/calibri.ttf',
            'bold': 'C:/Windows/Fonts/calibrib.ttf',
            'italic': 'C:/Windows/Fonts/calibrii.ttf',
            'bolditalic': 'C:/Windows/Fonts/calibriz.ttf'
        },
        'urls': {}
    },
    'segoeui': {
        'id': 'segoeui',
        'name': 'Segoe UI',
        'family': 'Segoe UI',
        'prefix': 'SegoeUI_',
        'badge': 'System Font',
        'description': 'Font antarmuka modern Windows Fluent Design dengan tipografi tajam.',
        'files': {
            'regular': 'C:/Windows/Fonts/segoeui.ttf',
            'bold': 'C:/Windows/Fonts/segoeuib.ttf',
            'italic': 'C:/Windows/Fonts/segoeuii.ttf',
            'bolditalic': 'C:/Windows/Fonts/segoeuiz.ttf'
        },
        'urls': {}
    }
}

ROBOTO_FONTS = FONTS_REGISTRY['roboto']['files']

def get_target_font(font_key: Optional[str] = None):
    raw_key = (font_key or 'roboto').strip().lower().replace('-', '').replace(' ', '').replace('_', '')
    if raw_key in ['auto', 'default', '']:
        return FONTS_REGISTRY['roboto']
    for reg_key, cfg in FONTS_REGISTRY.items():
        clean_name = cfg['name'].lower().replace(' ', '').replace('-', '').replace('_', '')
        clean_fam = cfg['family'].lower().replace(' ', '').replace('-', '').replace('_', '')
        if raw_key == reg_key or raw_key == clean_name or raw_key == clean_fam:
            return cfg
    return FONTS_REGISTRY['roboto']

def ensure_all_fonts():
    for f_cfg in FONTS_REGISTRY.values():
        for filename, url in f_cfg.get('urls', {}).items():
            dest = os.path.join(FONTS_DIR, filename)
            if not os.path.exists(dest) or os.path.getsize(dest) == 0:
                try:
                    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
                    with urllib.request.urlopen(req, timeout=12) as resp:
                        data = resp.read()
                        with open(dest, 'wb') as f:
                            f.write(data)
                except Exception as e:
                    print(f"Warning: Gagal mengunduh {filename}: {e}")

ensure_all_fonts()

def ensure_roboto_fonts():
    ensure_all_fonts()

def convert_xml_to_font(xml_bytes: bytes, target_font: str = "roboto") -> tuple[bytes, int]:
    font_cfg = get_target_font(target_font)
    target_family = font_cfg['family']
    count = 0
    try:
        xml_str = xml_bytes.decode('utf-8')
    except UnicodeDecodeError:
        xml_str = xml_bytes.decode('latin-1', errors='replace')

    font_attr_names = {
        'family', 'font-family', 'fontfamily', 'font_family',
        'face', 'typeface', 'font', 'font-name', 'fontname', 'font_name'
    }

    known_source_fonts = [
        'ArialMT', 'Arial-BoldMT', 'Arial-ItalicMT', 'Arial-BoldItalicMT',
        'Arial', 'Calibri', 'Times New Roman', 'TimesNewRoman', 'Times',
        'Neo Sans Pro', 'NeoSansPro', 'Neo Sans Intel', 'NeoSans Intel',
        'Neo-Sans-Pro', 'Neo-Sans', 'Neo Sans', 'NeoSans',
        'Segoe UI', 'SegoeUI', 'Helvetica', 'Tahoma', 'Verdana',
        'Aptos', 'Cambria', 'Garamond'
    ]

    try:
        root = ET.fromstring(xml_str)
        for elem in root.iter():
            for key, val in list(elem.attrib.items()):
                clean_key = key.lower().split('}')[-1]
                if clean_key in font_attr_names:
                    if val != target_family:
                        elem.attrib[key] = target_family
                        count += 1
                elif val:
                    modified = False
                    for src in known_source_fonts:
                        if src.lower() in val.lower():
                            val = re.sub(rf'\b{re.escape(src)}\b', target_family, val, flags=re.IGNORECASE)
                            modified = True
                    if modified:
                        elem.attrib[key] = val
                        count += 1

            if elem.text:
                modified = False
                new_text = elem.text
                for src in known_source_fonts:
                    if src.lower() in new_text.lower():
                        new_text = re.sub(rf'\b{re.escape(src)}\b', target_family, new_text, flags=re.IGNORECASE)
                        modified = True
                if modified:
                    elem.text = new_text
                    count += 1

            if elem.tail:
                modified = False
                new_tail = elem.tail
                for src in known_source_fonts:
                    if src.lower() in new_tail.lower():
                        new_tail = re.sub(rf'\b{re.escape(src)}\b', target_family, new_tail, flags=re.IGNORECASE)
                        modified = True
                if modified:
                    elem.tail = new_tail
                    count += 1

        out_stream = io.BytesIO()
        tree = ET.ElementTree(root)
        tree.write(out_stream, encoding='utf-8', xml_declaration=True)
        converted_xml = out_stream.getvalue()

    except Exception:
        pattern = r'(?i)\b(Arial|Calibri|Times New Roman|Neo[\s\-_]?Sans(?:[\s\-_]?Pro|[\s\-_]?Intel)?|NeoSansPro|NeoSans|Segoe UI|Helvetica)\b'
        converted_str, count = re.subn(pattern, target_family, xml_str)
        converted_xml = converted_str.encode('utf-8')

    converted_str = converted_xml.decode('utf-8', errors='replace')
    extra_pattern = r'(?i)\b(Arial|Calibri|Times New Roman|Times|Neo[\s\-_]?Sans(?:[\s\-_]?Pro|[\s\-_]?Intel)?|NeoSansPro|NeoSans)\b'
    extra_subs, extra_count = re.subn(extra_pattern, target_family, converted_str)
    if extra_count > 0:
        converted_xml = extra_subs.encode('utf-8')
        count += extra_count

    return converted_xml, count

def convert_xml_arial_to_roboto(xml_bytes: bytes) -> tuple[bytes, int]:
    return convert_xml_to_font(xml_bytes, 'roboto')


def convert_rtf_to_font(rtf_bytes: bytes, target_font: str = "roboto") -> tuple[bytes, int]:
    font_cfg = get_target_font(target_font)
    target_family = font_cfg['family']
    count = 0

    encoding_used = 'latin-1'
    try:
        rtf_str = rtf_bytes.decode('utf-8')
        if any(ord(c) > 0xFF for c in rtf_str):
            encoding_used = 'utf-8'
        else:
            rtf_str = rtf_bytes.decode('latin-1')
            encoding_used = 'latin-1'
    except UnicodeDecodeError:
        try:
            rtf_str = rtf_bytes.decode('latin-1')
            encoding_used = 'latin-1'
        except UnicodeDecodeError:
            rtf_str = rtf_bytes.decode('cp1252', errors='replace')
            encoding_used = 'cp1252'

    if not rtf_str.strip().startswith('{\\rtf'):
        raise ValueError("File bukan dokumen RTF yang valid.")

    font_decl_pattern = re.compile(
        r'(\\f\d+[^;{]*?)\s+([^;{]+)\s*;',
        re.IGNORECASE
    )

    def replace_in_fonttbl(m):
        nonlocal count
        original_name = m.group(2)
        lower = original_name.lower()
        if 'bold' in lower and 'italic' in lower:
            new_name = f"{target_family} Bold Italic"
        elif 'boldmt' in lower or 'bold' in lower:
            new_name = f"{target_family} Bold"
        elif 'italic' in lower or 'oblique' in lower:
            new_name = f"{target_family} Italic"
        elif 'condensed' in lower or 'narrow' in lower:
            new_name = f"{target_family} Condensed"
        else:
            new_name = target_family
        count += 1
        return f"{m.group(1)} {new_name};"

    fonttbl_pattern = re.compile(r'(\{\\fonttbl)(.*?)(\})', re.DOTALL)

    def process_fonttbl_block(m):
        prefix = m.group(1)
        body = font_decl_pattern.sub(replace_in_fonttbl, m.group(2))
        suffix = m.group(3)
        return prefix + body + suffix

    rtf_str = fonttbl_pattern.sub(process_fonttbl_block, rtf_str)

    src_pattern = r'\b(Arial|Calibri|Times New Roman|Segoe UI|Helvetica|Neo[\s\-_]?Sans(?:[\s\-_]?Pro|[\s\-_]?Intel)?|NeoSansPro|NeoSans)\b'
    rtf_str, extra_subs = re.subn(src_pattern, target_family, rtf_str, flags=re.IGNORECASE)
    count += extra_subs

    result_bytes = rtf_str.encode(encoding_used, errors='replace')
    return result_bytes, count

def convert_rtf_arial_to_roboto(rtf_bytes: bytes) -> tuple[bytes, int]:
    return convert_rtf_to_font(rtf_bytes, 'roboto')


def convert_doc_to_font(doc_bytes: bytes, target_font: str = "roboto") -> tuple[bytes, int]:
    for encoding in ('utf-8', 'latin-1', 'cp1252'):
        try:
            peek = doc_bytes[:20].decode(encoding)
            if peek.strip().startswith('{\\rtf'):
                return convert_rtf_to_font(doc_bytes, target_font)
            break
        except UnicodeDecodeError:
            continue

    try:
        converted_bytes, count = convert_docx_to_font(doc_bytes, target_font)
        return converted_bytes, count
    except Exception:
        pass

    font_cfg = get_target_font(target_font)
    target_family = font_cfg['family']
    count = 0

    target_5 = target_family[:5].encode('latin-1', errors='replace').ljust(5)

    arial_variants = [
        (b'Arial-BoldItalicMT', (target_family + '-BoldItalic').encode('latin-1', errors='replace').ljust(18)[:18]),
        (b'Arial-BoldMT',       (target_family + '-Bold').encode('latin-1', errors='replace').ljust(12)[:12]),
        (b'Arial-ItalicMT',     (target_family + '-Italic').encode('latin-1', errors='replace').ljust(14)[:14]),
        (b'ArialMT',            target_family.encode('latin-1', errors='replace').ljust(7)[:7]),
        (b'Arial Unicode MS',   target_family.encode('latin-1', errors='replace').ljust(16)[:16]),
        (b'Arial Narrow',       (target_family + ' Cond.').encode('latin-1', errors='replace').ljust(12)[:12]),
        (b'Arial Bold',         (target_family + ' Bold').encode('latin-1', errors='replace').ljust(10)[:10]),
        (b'Arial Italic',       (target_family + ' Italic').encode('latin-1', errors='replace').ljust(12)[:12]),
    ]

    result = doc_bytes
    for old, new in arial_variants:
        if len(old) != len(new):
            continue
        n = result.count(old)
        if n > 0:
            result = result.replace(old, new)
            count += n

    arial_bin_pat = re.compile(b'(?<=\x00)Arial(?=[\x00\x20])', re.IGNORECASE)
    result, n = arial_bin_pat.subn(target_5, result)
    count += n

    arial_wide = 'Arial'.encode('utf-16-le')
    target_wide = target_5.decode('latin-1', errors='replace').encode('utf-16-le')
    n_wide = result.count(arial_wide)
    if n_wide > 0:
        result = result.replace(arial_wide, target_wide)
        count += n_wide

    return result, count

def convert_doc_arial_to_roboto(doc_bytes: bytes) -> tuple[bytes, int]:
    return convert_doc_to_font(doc_bytes, 'roboto')


def convert_docx_to_font(docx_bytes: bytes, target_font: str = "roboto") -> tuple[bytes, int]:
    font_cfg = get_target_font(target_font)
    target_family = font_cfg['family']
    doc = docx.Document(io.BytesIO(docx_bytes))
    count = 0

    def process_run(run):
        nonlocal count
        changed = False

        if run.font.name != target_family:
            run.font.name = target_family
            changed = True

        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is not None:
            for attr in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
                val = rFonts.get(qn(f'w:{attr}'))
                if val != target_family:
                    rFonts.set(qn(f'w:{attr}'), target_family)
                    changed = True
        else:
            rFonts_elem = docx.oxml.OxmlElement('w:rFonts')
            rFonts_elem.set(qn('w:ascii'), target_family)
            rFonts_elem.set(qn('w:hAnsi'), target_family)
            rFonts_elem.set(qn('w:cs'), target_family)
            rPr.append(rFonts_elem)
            changed = True

        if changed:
            count += 1

    for style in doc.styles:
        try:
            if hasattr(style, 'font') and style.font is not None:
                if style.font.name != target_family:
                    style.font.name = target_family
                    count += 1
        except Exception:
            pass

    for p in doc.paragraphs:
        for run in p.runs:
            process_run(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        process_run(run)

    for section in doc.sections:
        for header_type in [section.header, section.first_page_header, section.even_page_header]:
            if header_type is not None:
                for p in header_type.paragraphs:
                    for run in p.runs:
                        process_run(run)
                for table in header_type.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    process_run(run)

        for footer_type in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer_type is not None:
                for p in footer_type.paragraphs:
                    for run in p.runs:
                        process_run(run)
                for table in footer_type.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    process_run(run)

    out_stream = io.BytesIO()
    doc.save(out_stream)
    return out_stream.getvalue(), count

def convert_docx_arial_to_roboto(docx_bytes: bytes) -> tuple[bytes, int]:
    return convert_docx_to_font(docx_bytes, 'roboto')


ROUND_BULLETS = {
    '\uf0b7', '\uf06c', '\uf06d', '\u25cf', '\u25cb', '\u25ef',
    '\u2022', '\u00b7', '\u2219', '\u2023'
}
SQUARE_BULLETS = {
    '\uf0a7', '\uf06e', '\uf071', '\uf0de', '\uf0d8', '\uf0a8',
    '\u25aa', '\u25ab', '\u25a0', '\u25a1', '\u25fe', '\u25fc',
    '\u25c6', '\u25c7', '\xa7', '\x00'
}
ALL_BULLETS = ROUND_BULLETS | SQUARE_BULLETS | {
    '\u00bb', '\u00ab', '\u2013', '\u2014', '\u25b8', '\u25b9',
    '\u25ba', '\u25bb', '\u25c0', '\u25c1', '\u25c2', '\u25c3',
    '-', '*', '>'
}

PDF_DELIMITERS = set(' \t\r\n\x00\x0c()<>[]{}/%')

def _is_square_bullet(text: str, font_str: str = '') -> bool:
    if text in SQUARE_BULLETS:
        return True
    if len(text) == 1:
        code = ord(text)
        if 0xF000 <= code <= 0xF0FF and text not in ROUND_BULLETS:
            return True
        if 'wingdings' in font_str and (text in ['n', '\xa7', 'q'] or text not in ROUND_BULLETS):
            return True
    return False

def _is_bullet_span(text: str, font_str: str = '') -> bool:
    t = text.strip()
    if not t:
        return False
    if t in ALL_BULLETS:
        return True
    if _is_square_bullet(t, font_str):
        return True
    if len(t) == 1 and not t.isalnum():
        return True
    if re.match(r'^\(?[0-9a-zA-Z]{1,3}[\.\)]$', t):
        return True
    return False

def _remove_text_from_content_stream(stream_text: str) -> str:
    out = []
    i = 0
    n = len(stream_text)
    in_text = False
    bt_non_str = []
    while i < n:
        if not in_text:
            if (i == 0 or stream_text[i-1] in PDF_DELIMITERS) and stream_text[i:i+2] == 'BT' and (i+2 == n or stream_text[i+2] in PDF_DELIMITERS):
                in_text = True
                bt_non_str = []
                i += 2
                continue
            out.append(stream_text[i])
            i += 1
        else:
            if stream_text[i] == '(':
                i += 1
                depth = 1
                while i < n and depth > 0:
                    if stream_text[i] == '\\':
                        i += 2
                    elif stream_text[i] == '(':
                        depth += 1
                        i += 1
                    elif stream_text[i] == ')':
                        depth -= 1
                        i += 1
                    else:
                        i += 1
                continue
            elif stream_text[i] == '<' and (i+1 < n and stream_text[i+1] != '<'):
                i += 1
                while i < n and stream_text[i] != '>':
                    i += 1
                if i < n:
                    i += 1
                continue
            elif (i == 0 or stream_text[i-1] in PDF_DELIMITERS) and stream_text[i:i+2] == 'ET' and (i+2 == n or stream_text[i+2] in PDF_DELIMITERS):
                in_text = False
                non_str_content = ''.join(bt_non_str)
                state_ops = re.findall(r'(/[\w\#]+)\s+(cs|CS|gs)\b', non_str_content)
                if state_ops:
                    out.append(' ' + ' '.join(f'{name} {op}' for name, op in state_ops) + ' ')
                color_ops = re.findall(r'((?:[0-9\.\+\-]+\s+){1,4}(?:k|K|rg|RG|g|G))\b', non_str_content)
                if color_ops:
                    out.append(' ' + ' '.join(color_ops) + ' ')
                i += 2
                continue
            else:
                bt_non_str.append(stream_text[i])
                i += 1
    return ''.join(out)

def _clear_font_dict(obj_str: str) -> str:
    m = re.search(r'/Font\s*<<', obj_str)
    if not m:
        return obj_str
    dict_start = m.end() - 2
    i = dict_start + 2
    depth = 1
    n = len(obj_str)
    while i < n and depth > 0:
        if obj_str[i:i+2] == '<<':
            depth += 1
            i += 2
        elif obj_str[i:i+2] == '>>':
            depth -= 1
            i += 2
        else:
            i += 1
    if depth == 0:
        return obj_str[:dict_start] + '<<>>' + obj_str[i:]
    return obj_str

def _clear_all_font_refs(obj_str: str) -> str:
    result = re.sub(r'/Font\s+\d+\s+0\s+R\b', '/Font <<>>', obj_str)
    result = _clear_font_dict(result)
    return result

def convert_pdf_to_font(pdf_bytes: bytes, target_font: str = "roboto") -> tuple[bytes, int]:
    font_cfg = get_target_font(target_font)
    family_clean = font_cfg['family'].replace(' ', '')
    fallback_cfg = FONTS_REGISTRY['roboto']

    reg_path = font_cfg['files'].get('regular', '')
    bold_path = font_cfg['files'].get('bold', '')
    italic_path = font_cfg['files'].get('italic', '')
    bolditalic_path = font_cfg['files'].get('bolditalic', '')

    if not (reg_path and os.path.exists(reg_path)):
        reg_path = fallback_cfg['files']['regular']
    if not (bold_path and os.path.exists(bold_path)):
        bold_path = fallback_cfg['files']['bold']
    if not (italic_path and os.path.exists(italic_path)):
        italic_path = fallback_cfg['files']['italic']
    if not (bolditalic_path and os.path.exists(bolditalic_path)):
        bolditalic_path = fallback_cfg['files']['bolditalic']

    font_cache = {}
    def get_font_obj(path):
        if path and path not in font_cache and os.path.exists(path):
            font_cache[path] = fitz.Font(fontfile=path)
        return font_cache.get(path)

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    total_spans_converted = 0

    for page in doc:
        try:
            page.clean_contents(sanitize=False)
        except Exception:
            pass

    all_pages_elements = []
    for page in doc:
        d = page.get_text("dict")
        raw_spans = []
        for b in d.get("blocks", []):
            if b.get("type") != 0:
                continue
            for line in b.get("lines", []):
                for s in line.get("spans", []):
                    if s.get("text", "").strip():
                        raw_spans.append(s)

        superseded_ids = set()
        for i, s1 in enumerate(raw_spans):
            t1 = s1.get("text", "").strip().lower()
            r1 = fitz.Rect(s1["bbox"])
            for s2 in raw_spans[i+1:]:
                t2 = s2.get("text", "").strip().lower()
                if t1 == t2 or (t1 in t2 and len(t1) > 5) or (t2 in t1 and len(t2) > 5):
                    r2 = fitz.Rect(s2["bbox"])
                    if r1.intersects(r2):
                        intersect_area = fitz.Rect(r1).intersect(r2).get_area()
                        min_area = min(r1.get_area(), r2.get_area())
                        if min_area > 0 and (intersect_area / min_area) > 0.4:
                            superseded_ids.add(id(s1))
                            break

        page_pix = page.get_pixmap(dpi=72)
        pix_w, pix_h = page_pix.width, page_pix.height
        pix_samples = page_pix.samples
        for s in raw_spans:
            if id(s) in superseded_ids:
                continue
            bx0 = max(0, min(pix_w - 1, int(s["bbox"][0])))
            by0 = max(0, min(pix_h - 1, int(s["bbox"][1])))
            bx1 = max(0, min(pix_w - 1, int(s["bbox"][2])))
            by1 = max(0, min(pix_h - 1, int(s["bbox"][3])))
            if bx1 > bx0 and by1 > by0:
                s_cols = set()
                for py in range(by0, by1 + 1):
                    for px in range(bx0, bx1 + 1):
                        p_idx = (py * pix_w + px) * 3
                        s_cols.add((pix_samples[p_idx], pix_samples[p_idx+1], pix_samples[p_idx+2]))
                        if len(s_cols) > 2:
                            break
                    if len(s_cols) > 2:
                        break
                if len(s_cols) <= 2:
                    superseded_ids.add(id(s))

        x1_counts = {}
        for b in d.get("blocks", []):
            if b.get("type") != 0:
                continue
            for l in b.get("lines", []):
                txt = " ".join(s.get("text", "") for s in l.get("spans", []) if id(s) not in superseded_ids).strip()
                if len(txt.split()) >= 3 and not txt.endswith(":"):
                    x1_val = round(l["bbox"][2], 0)
                    x1_counts[x1_val] = x1_counts.get(x1_val, 0) + 1
        common_margins = [k for k, v in x1_counts.items() if v >= 4]

        all_page_groups = []
        for b in d.get("blocks", []):
            if b.get("type") != 0:
                continue
            for l in b.get("lines", []):
                sps = [s for s in l.get("spans", []) if s.get("text", "").strip() and id(s) not in superseded_ids]
                if sps:
                    sps.sort(key=lambda s: s["origin"][0])
                    all_page_groups.append({
                        "baseline_y": round(sps[0]["origin"][1], 1),
                        "bbox": tuple(l["bbox"]),
                        "spans": sps
                    })

        all_page_groups.sort(key=lambda g: (g["baseline_y"], g["bbox"][0]))

        page_bullet_indents = {}
        for gi, g in enumerate(all_page_groups):
            sps = g["spans"]
            t0 = sps[0].get("text", "").strip()
            if _is_bullet_span(t0, sps[0].get("font", "")) and len(sps) >= 2:
                p_key = round(sps[0]["origin"][0], 0)
                for next_g in all_page_groups[gi+1:]:
                    nt0 = next_g["spans"][0].get("text", "").strip()
                    if _is_bullet_span(nt0, next_g["spans"][0].get("font", "")):
                        break
                    cont_x0 = round(next_g["bbox"][0], 1)
                    if cont_x0 > sps[0]["bbox"][2] + 1.0:
                        page_bullet_indents.setdefault(p_key, []).append(cont_x0)
                        break

        page_std_indent = {}
        for p_key, c_list in page_bullet_indents.items():
            if c_list:
                freq = {}
                for cv in c_list:
                    freq[cv] = freq.get(cv, 0) + 1
                page_std_indent[p_key] = max(freq.items(), key=lambda item: item[1])[0]

        elements_to_draw = []
        for b in d.get("blocks", []):
            if b.get("type") != 0:
                continue
            lines = b.get("lines", [])
            if not lines:
                continue

            grouped_lines = []
            for l in lines:
                spans = [s for s in l.get("spans", []) if s.get("text", "").strip() and id(s) not in superseded_ids]
                if not spans:
                    continue
                base_y = round(spans[0]["origin"][1], 1)
                merged = False
                for g in grouped_lines:
                    if abs(g["baseline_y"] - base_y) <= 1.2:
                        h_gap = max(0, max(g["bbox"][0], l["bbox"][0]) - min(g["bbox"][2], l["bbox"][2]))
                        if h_gap <= 15.0:
                            g["lines"].append(l)
                            g["spans"].extend(spans)
                            g["bbox"] = (
                                min(g["bbox"][0], l["bbox"][0]),
                                min(g["bbox"][1], l["bbox"][1]),
                                max(g["bbox"][2], l["bbox"][2]),
                                max(g["bbox"][3], l["bbox"][3])
                            )
                            merged = True
                            break
                if not merged:
                    grouped_lines.append({
                        "baseline_y": base_y,
                        "lines": [l],
                        "spans": list(spans),
                        "bbox": tuple(l["bbox"])
                    })

            for g in grouped_lines:
                g["spans"].sort(key=lambda s: s["origin"][0])

            if not grouped_lines:
                continue

            block_max_x1 = max(g["bbox"][2] for g in grouped_lines)
            aligned_right_count = sum(1 for g in grouped_lines if abs(g["bbox"][2] - block_max_x1) <= 2.5)

            for gi, g in enumerate(grouped_lines):
                spans = g["spans"]
                if not spans:
                    continue

                prefix_span = None
                body_spans = spans
                if len(spans) >= 2:
                    s0 = spans[0]
                    t0 = s0.get("text", "").strip()
                    is_bullet = _is_bullet_span(t0, s0.get("font", ""))
                    if is_bullet:
                        bi = 1
                        while bi < len(spans) and not spans[bi].get("text", "").strip():
                            bi += 1
                        if bi < len(spans) and (spans[bi]["origin"][0] - s0["bbox"][2]) > 2.0:
                            prefix_span = s0
                            body_spans = spans[bi:]

                orig_body_x0 = body_spans[0]["origin"][0] if body_spans else g["bbox"][0]
                body_x0 = orig_body_x0

                if prefix_span and body_spans:
                    p_key = round(prefix_span["origin"][0], 0)
                    cand_indent = None
                    for next_g in grouped_lines[gi+1:]:
                        n_sps = next_g["spans"]
                        if not n_sps:
                            continue
                        nt0 = n_sps[0].get("text", "").strip()
                        if _is_bullet_span(nt0, n_sps[0].get("font", "")):
                            break
                        c_x0 = round(next_g["bbox"][0], 1)
                        if c_x0 > prefix_span["bbox"][2] + 1.0:
                            cand_indent = c_x0
                            break
                    if cand_indent is None:
                        cand_indent = page_std_indent.get(p_key)

                    if cand_indent is not None and cand_indent > prefix_span["bbox"][2] + 1.0:
                        if abs(cand_indent - orig_body_x0) <= 12.0:
                            body_x0 = cand_indent

                target_x1 = g["bbox"][2]
                target_w = target_x1 - body_x0
                is_right_aligned = (aligned_right_count >= 2 and abs(target_x1 - block_max_x1) <= 2.5) or any(abs(target_x1 - cm) <= 2.5 for cm in common_margins)


                words_data = []
                for s in body_spans:
                    stext = s.get("text", "")
                    if not stext:
                        continue
                    sz = s.get("size", 10.0)
                    c = s.get("color", 0)
                    cr, cg, cb = (((c >> 16) & 255) / 255.0, ((c >> 8) & 255) / 255.0, (c & 255) / 255.0)
                    color = (0.0, 0.0, 0.0) if (cr < 0.22 and cg < 0.22 and cb < 0.22) else (cr, cg, cb)
                    flags = s.get("flags", 0)
                    font_str = s.get("font", "").lower()
                    is_bold = bool(flags & 16) or any(w in font_str for w in ['bold', 'medium', 'semibold', 'semi-bold', 'demi', 'black', 'heavy'])
                    is_italic = bool(flags & 2) or 'italic' in font_str or 'oblique' in font_str
                    if is_bold and is_italic and os.path.exists(bolditalic_path):
                        fn, ff = 'Roboto-BoldItalic', bolditalic_path
                    elif is_bold and os.path.exists(bold_path):
                        fn, ff = 'Roboto-Bold', bold_path
                    elif is_italic and os.path.exists(italic_path):
                        fn, ff = 'Roboto-Italic', italic_path
                    elif os.path.exists(reg_path):
                        fn, ff = 'Roboto-Regular', reg_path
                    else:
                        fn, ff = 'helv', None
                    sf_obj = get_font_obj(ff)
                    for rb in ROUND_BULLETS:
                        stext = stext.replace(rb, '\u2022')
                    words = [w for w in stext.split() if w]
                    for w in words:
                        w_len = sf_obj.text_length(w, fontsize=sz) if sf_obj else sz * 0.5 * len(w)
                        words_data.append({
                            'word': w, 'fn': fn, 'ff': ff, 'sz': sz, 'color': color,
                            'w': w_len, 'y': s['origin'][1], 'f_obj': sf_obj
                        })

                has_square = any(_is_square_bullet(s.get("text", "").strip(), s.get("font", "").lower()) for s in body_spans)
                total_words = len(words_data)
                total_words_w = sum(wd['w'] for wd in words_data)
                default_sp = words_data[0]['f_obj'].text_length(' ', fontsize=words_data[0]['sz']) if words_data else 2.0


                body_scale = 1.0
                if total_words >= 2 and target_w > 0:
                    gaps = total_words - 1
                    natural_w = total_words_w + gaps * default_sp
                    if natural_w > target_w * 1.005:
                        body_scale = max(0.85, target_w / natural_w)
                        for wd in words_data:
                            wd['sz'] *= body_scale
                            wd['w'] = wd['f_obj'].text_length(wd['word'], fontsize=wd['sz']) if wd['f_obj'] else wd['sz'] * 0.5 * len(wd['word'])
                        total_words_w = sum(wd['w'] for wd in words_data)
                        default_sp = words_data[0]['f_obj'].text_length(' ', fontsize=words_data[0]['sz']) if words_data else 2.0

                full_line_txt = " ".join(wd['word'] for wd in words_data).strip()
                is_label = full_line_txt.endswith(":") or full_line_txt.endswith("?")
                is_centered_block = False
                if len(grouped_lines) >= 2:
                    centers = [(g_line["bbox"][0] + g_line["bbox"][2]) / 2.0 for g_line in grouped_lines]
                    widths = [g_line["bbox"][2] - g_line["bbox"][0] for g_line in grouped_lines]
                    if (max(centers) - min(centers) <= 2.5) and (max(widths) - min(widths) >= 4.0):
                        is_centered_block = True
                can_justify = False
                if is_right_aligned and not has_square and not is_label and not is_centered_block and total_words >= 4 and target_w > total_words_w:
                    gaps = total_words - 1
                    gap_w = (target_w - total_words_w) / gaps
                    if default_sp * 0.5 <= gap_w <= default_sp * 12.0:
                        can_justify = True

                def add_prefix_span(p_span):
                    p_txt = p_span.get("text", "").strip()
                    p_orig = fitz.Point(p_span["origin"])
                    p_sz = p_span.get("size", 10.0)
                    p_c = p_span.get("color", 0)
                    p_cr, p_cg, p_cb = (((p_c >> 16) & 255) / 255.0, ((p_c >> 8) & 255) / 255.0, (p_c & 255) / 255.0)
                    p_col = (0.0, 0.0, 0.0) if (p_cr < 0.22 and p_cg < 0.22 and p_cb < 0.22) else (p_cr, p_cg, p_cb)
                    p_flags = p_span.get("flags", 0)
                    p_fstr = p_span.get("font", "").lower()
                    if _is_square_bullet(p_txt, p_fstr):
                        bw = p_sz * 0.38
                        bh = p_sz * 0.38
                        r = fitz.Rect(p_orig.x, p_orig.y - p_sz * 0.48, p_orig.x + bw, p_orig.y - p_sz * 0.48 + bh)
                        elements_to_draw.append(('rect', r, p_col))
                    else:
                        for rb in ROUND_BULLETS:
                            p_txt = p_txt.replace(rb, '\u2022')
                        p_bold = bool(p_flags & 16) or any(w in p_fstr for w in ['bold', 'medium', 'semibold', 'semi-bold', 'demi', 'black', 'heavy'])
                        p_italic = bool(p_flags & 2) or 'italic' in p_fstr or 'oblique' in p_fstr
                        if p_bold and p_italic and os.path.exists(bolditalic_path):
                            p_fn, p_ff = 'Roboto-BoldItalic', bolditalic_path
                        elif p_bold and os.path.exists(bold_path):
                            p_fn, p_ff = 'Roboto-Bold', bold_path
                        elif p_italic and os.path.exists(italic_path):
                            p_fn, p_ff = 'Roboto-Italic', italic_path
                        elif os.path.exists(reg_path):
                            p_fn, p_ff = 'Roboto-Regular', reg_path
                        else:
                            p_fn, p_ff = 'helv', None
                        elements_to_draw.append(('text', p_orig, p_txt, p_fn, p_ff, p_sz, p_col))

                if can_justify:
                    if prefix_span:
                        add_prefix_span(prefix_span)
                    curr_x = body_x0
                    gaps = total_words - 1
                    gap_w = (target_w - total_words_w) / gaps
                    for i, wd in enumerate(words_data):
                        elements_to_draw.append(('text', fitz.Point(curr_x, wd['y']), wd['word'], wd['fn'], wd['ff'], wd['sz'], wd['color']))
                        if i < gaps:
                            gap_start = curr_x + wd['w']
                            sp_x = gap_start + (gap_w - default_sp) / 2.0
                            elements_to_draw.append(('text', fitz.Point(sp_x, wd['y']), ' ', wd['fn'], wd['ff'], wd['sz'], wd['color']))
                        curr_x += wd['w'] + gap_w
                    continue


                if prefix_span:
                    add_prefix_span(prefix_span)
                    draw_spans = body_spans
                    avail_w = target_w
                    base_start_x = body_x0
                else:
                    draw_spans = spans
                    avail_w = g["bbox"][2] - g["bbox"][0]
                    base_start_x = g["bbox"][0]

                rob_body_w = 0.0
                for s in draw_spans:
                    stext = s.get("text", "")
                    if not stext:
                        continue
                    for rb in ROUND_BULLETS:
                        stext = stext.replace(rb, '\u2022')
                    s_sz = s.get("size", 10.0)
                    s_flags = s.get("flags", 0)
                    s_fstr = s.get("font", "").lower()
                    s_bold = bool(s_flags & 16) or any(w in s_fstr for w in ['bold', 'medium', 'semibold', 'semi-bold', 'demi', 'black', 'heavy'])
                    s_italic = bool(s_flags & 2) or 'italic' in s_fstr or 'oblique' in s_fstr
                    if s_bold and s_italic and os.path.exists(bolditalic_path):
                        s_ff = bolditalic_path
                    elif s_bold and os.path.exists(bold_path):
                        s_ff = bold_path
                    elif s_italic and os.path.exists(italic_path):
                        s_ff = italic_path
                    elif os.path.exists(reg_path):
                        s_ff = reg_path
                    else:
                        s_ff = None
                    sf_obj = get_font_obj(s_ff)
                    clean_st = stext.strip()
                    if clean_st and _is_square_bullet(clean_st, s_fstr):
                        rob_body_w += s_sz * 0.38
                    else:
                        rob_body_w += sf_obj.text_length(stext, fontsize=s_sz) if sf_obj else s_sz * 0.5 * len(stext)

                line_scale = 1.0
                if avail_w > 0 and rob_body_w > avail_w * 1.005:
                    line_scale = max(0.85, avail_w / rob_body_w)

                body_shift_x = (body_x0 - orig_body_x0) if prefix_span else 0.0
                prev_end_x = 0
                prev_orig_x1 = 0
                for s in draw_spans:
                    text = s.get("text", "")
                    if not text:
                        continue
                    origin = fitz.Point(s["origin"][0] + body_shift_x, s["origin"][1])
                    size = s["size"] * line_scale
                    c = s.get("color", 0)
                    cr, cg, cb = (((c >> 16) & 255) / 255.0, ((c >> 8) & 255) / 255.0, (c & 255) / 255.0)
                    color = (0.0, 0.0, 0.0) if (cr < 0.22 and cg < 0.22 and cb < 0.22) else (cr, cg, cb)
                    flags = s.get("flags", 0)
                    font_str = s.get("font", "").lower()
                    is_bold = bool(flags & 16) or any(w in font_str for w in ['bold', 'medium', 'semibold', 'semi-bold', 'demi', 'black', 'heavy'])
                    is_italic = bool(flags & 2) or 'italic' in font_str or 'oblique' in font_str

                    if is_bold and is_italic and os.path.exists(bolditalic_path):
                        fn, ff = f'{family_clean}-BoldItalic', bolditalic_path
                    elif is_bold and os.path.exists(bold_path):
                        fn, ff = f'{family_clean}-Bold', bold_path
                    elif is_italic and os.path.exists(italic_path):
                        fn, ff = f'{family_clean}-Italic', italic_path
                    elif os.path.exists(reg_path):
                        fn, ff = f'{family_clean}-Regular', reg_path
                    else:
                        fn, ff = 'helv', None

                    f_obj = get_font_obj(ff)
                    for rb in ROUND_BULLETS:
                        text = text.replace(rb, '\u2022')

                    clean_t = text.strip()
                    if clean_t and _is_square_bullet(clean_t, font_str):
                        bw = size * 0.38
                        bh = size * 0.38
                        r = fitz.Rect(origin.x, origin.y - size * 0.48, origin.x + bw, origin.y - size * 0.48 + bh)
                        elements_to_draw.append(('rect', r, color))
                        prev_end_x = origin.x + bw
                        prev_orig_x1 = s["bbox"][2]
                        continue

                    if clean_t and any(_is_square_bullet(ch, font_str) for ch in text):
                        curr_x = origin.x
                        curr_y = origin.y
                        buf = ''
                        for ch in text:
                            if _is_square_bullet(ch, font_str):
                                if buf:
                                    elements_to_draw.append(('text', fitz.Point(curr_x, curr_y), buf, fn, ff, size, color))
                                    curr_x += f_obj.text_length(buf, fontsize=size) if f_obj else size * 0.5 * len(buf)
                                    buf = ''
                                bw = size * 0.38
                                bh = size * 0.38
                                r = fitz.Rect(curr_x, curr_y - size * 0.48, curr_x + bw, curr_y - size * 0.48 + bh)
                                elements_to_draw.append(('rect', r, color))
                                curr_x += bw + size * 0.15
                            else:
                                buf += ch
                        if buf:
                            elements_to_draw.append(('text', fitz.Point(curr_x, curr_y), buf, fn, ff, size, color))
                            curr_x += f_obj.text_length(buf, fontsize=size) if f_obj else size * 0.5 * len(buf)
                        prev_end_x = curr_x
                        prev_orig_x1 = s["bbox"][2]
                        continue

                    sp_char_w = f_obj.text_length(' ', fontsize=size) if f_obj else size * 0.25
                    orig_gap = (s["bbox"][0] - prev_orig_x1) if prev_orig_x1 > 0 else 999.0
                    if prev_end_x > 0:
                        if orig_gap <= 0.5:
                            origin = fitz.Point(prev_end_x, origin.y)
                        elif origin.x < prev_end_x:
                            origin = fitz.Point(prev_end_x, origin.y)
                        elif origin.x > prev_end_x + sp_char_w * 1.2:
                            span_gap = origin.x - prev_end_x
                            sp_x = prev_end_x + (span_gap - sp_char_w) / 2.0
                            elements_to_draw.append(('text', fitz.Point(sp_x, origin.y), ' ', fn, ff, size, color))

                    span_w = f_obj.text_length(text, fontsize=size) if f_obj else size * 0.5 * len(text)
                    prev_end_x = origin.x + span_w
                    prev_orig_x1 = s["bbox"][2]
                    elements_to_draw.append(('text', origin, text, fn, ff, size, color))

        all_pages_elements.append(elements_to_draw)


    _gfont_indirect_re = re.compile(r'/Font\s+\d+\s+0\s+R\b')
    for _gx in range(1, doc.xref_length()):
        try:
            _gobj = doc.xref_object(_gx, compressed=False)
            if not _gobj or '/Font' not in _gobj:
                continue
            _gnew = _gfont_indirect_re.sub('/Font <<>>', _gobj)
            _gnew = _clear_font_dict(_gnew)
            if _gnew != _gobj:
                doc.update_object(_gx, _gnew)
        except Exception:
            pass

    for xref in range(1, doc.xref_length()):
        try:
            obj_str = doc.xref_object(xref, compressed=False)
            if '/Subtype /Form' in obj_str or '/Subtype/Form' in obj_str:
                raw_bytes = doc.xref_stream(xref)
                if raw_bytes and b'BT' in raw_bytes:
                    m = re.search(r'/Resources\s+(\d+)\s+0\s+R', obj_str)
                    if m:
                        res_xref = int(m.group(1))
                        res_obj = doc.xref_object(res_xref, compressed=False)
                        new_res = _clear_all_font_refs(res_obj)
                        if new_res != res_obj:
                            doc.update_object(res_xref, new_res)
                    elif '/Font' in obj_str:
                        new_obj = _clear_all_font_refs(obj_str)
                        if new_obj != obj_str:
                            doc.update_object(xref, new_obj)

                    raw_data = raw_bytes.decode('latin-1', errors='replace')
                    clean_data = _remove_text_from_content_stream(raw_data)
                    doc.update_stream(xref, clean_data.encode('latin-1'))
        except Exception:
            pass


    for page_idx, page in enumerate(doc):
        try:
            page.clean_contents(sanitize=False)
        except Exception:
            pass
        contents = page.get_contents()
        for cx in contents:
            try:
                raw_data = doc.xref_stream(cx).decode('latin-1', errors='replace')
                clean_data = _remove_text_from_content_stream(raw_data)
                doc.update_stream(cx, clean_data.encode('latin-1'))
            except Exception:
                pass

        try:
            p_obj = doc.xref_object(page.xref, compressed=False)
            m = re.search(r'/Resources\s+(\d+)\s+0\s+R', p_obj)
            if m:
                res_xref = int(m.group(1))
                res_obj = doc.xref_object(res_xref, compressed=False)
                new_res = _clear_all_font_refs(res_obj)
                if new_res != res_obj:
                    doc.update_object(res_xref, new_res)
            elif '/Font' in p_obj:
                new_p_obj = _clear_all_font_refs(p_obj)
                if new_p_obj != p_obj:
                    doc.update_object(page.xref, new_p_obj)
        except Exception:
            pass

        for item in all_pages_elements[page_idx]:
            try:
                if item[0] == 'rect':
                    page.draw_rect(item[1], color=None, fill=item[2], overlay=True)
                    total_spans_converted += 1
                elif item[0] == 'text':
                    _, orig, txt, fn, ff, sz, col = item
                    if ff:
                        page.insert_text(orig, txt, fontname=fn, fontfile=ff, fontsize=sz, color=col, overlay=True)
                    else:
                        page.insert_text(orig, txt, fontname=fn, fontsize=sz, color=col, overlay=True)
                    total_spans_converted += 1
            except Exception:
                continue

    for page in doc:
        try:
            page.clean_contents(sanitize=False)
        except Exception:
            pass

    try:
        doc.subset_fonts()
    except Exception:
        pass

    _fn_re = re.compile(r'/(?:BaseFont|FontName)\s*/([^\s/\[\]<>()\r\n]+)', re.IGNORECASE)
    _ftype_re = re.compile(r'/Type\s*/(?:Font|FontDescriptor)\b|/Subtype\s*/(?:Type0|Type1|TrueType|CIDFontType0|CIDFontType2|MMType1|Type3|OpenType)\b', re.IGNORECASE)
    for x in range(1, doc.xref_length()):
        try:
            obj_str = doc.xref_object(x, compressed=False)
            if not obj_str or obj_str.strip() in ('', 'null', '<< >>'):
                continue
            if not _ftype_re.search(obj_str):
                continue
            nm = _fn_re.search(obj_str)
            raw_name = nm.group(1).lower() if nm else ''
            clean_name = re.sub(r'^[A-Za-z0-9]{6}\+', '', raw_name)
            if 'roboto' in clean_name:
                continue
            try:
                s = doc.xref_stream(x)
                if s is not None:
                    doc.update_stream(x, b'\n')
            except Exception:
                pass
            try:
                doc.update_object(x, '<< >>')
            except Exception:
                pass
        except Exception:
            pass

    for x in range(1, doc.xref_length()):
        try:
            s_bytes = doc.xref_stream(x)
            if s_bytes and b'begincmap' in s_bytes:
                s = s_bytes.decode('latin-1', errors='replace')
                def fix_hex(m):
                    val = m.group(1)
                    if len(val) % 2 != 0:
                        if len(val) == 5:
                            cp = int(val, 16)
                            return '<' + chr(cp).encode('utf-16be').hex() + '>'
                        return '<0' + val + '>'
                    return '<' + val + '>'
                new_s = re.sub(r'<([0-9a-fA-F]+)>', fix_hex, s)
                if new_s != s:
                    doc.update_stream(x, new_s.encode('latin-1'))
        except Exception:
            pass

    out_pdf = doc.tobytes(garbage=4, deflate=True, clean=False)
    doc.close()
    return out_pdf, total_spans_converted

def convert_pdf_all_to_roboto(pdf_bytes: bytes) -> tuple[bytes, int]:
    return convert_pdf_to_font(pdf_bytes, 'roboto')


def load_dataset(file_bytes: bytes, filename: str) -> pd.DataFrame:
    ext = filename.lower().split('.')[-1]
    if ext == 'csv':
        try:
            df = pd.read_csv(io.BytesIO(file_bytes), dtype=str, keep_default_na=False, encoding='utf-8')
        except UnicodeDecodeError:
            df = pd.read_csv(io.BytesIO(file_bytes), dtype=str, keep_default_na=False, encoding='latin-1')
    elif ext in ['xlsx', 'xls']:
        df = pd.read_excel(io.BytesIO(file_bytes), dtype=str, keep_default_na=False)
    else:
        raise HTTPException(status_code=400, detail="Format file data nasabah harus berupa CSV atau Excel (.xlsx, .xls)")

    df.columns = [str(c).strip() for c in df.columns]
    return df


def get_document_clean_names(row: pd.Series, idx: int, columns: list[str]) -> tuple[str, str]:
    id_val = None
    for key in ['NO_POLIS', 'NOMOR_POLIS', 'POLIS', 'ID', 'NO']:
        for c in columns:
            if c.upper() == key:
                id_val = str(row[c]).strip()
                break
        if id_val:
            break

    name_val = None
    for key in ['NAMA', 'NAMA_NASABAH', 'PEMEGANG_POLIS']:
        for c in columns:
            if c.upper() == key:
                name_val = str(row[c]).strip()
                break
        if name_val:
            break

    clean_id = re.sub(r'[^a-zA-Z0-9_-]', '_', id_val) if id_val else f"NASABAH_{idx+1:04d}"
    clean_name = ("_" + re.sub(r'[^a-zA-Z0-9_-]', '_', name_val)) if name_val else ""
    return clean_id, clean_name


def generate_batch_xml_zip(template_bytes: bytes, df: pd.DataFrame) -> tuple[io.BytesIO, int]:
    try:
        template_str = template_bytes.decode('utf-8')
    except UnicodeDecodeError:
        template_str = template_bytes.decode('latin-1', errors='replace')

    zip_buffer = io.BytesIO()
    total_generated = 0

    with zipfile.ZipFile(zip_buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        for idx, row in df.iterrows():
            rendered_xml = template_str

            for col in df.columns:
                val = str(row[col]) if row[col] is not None else ""
                patterns = [
                    re.compile(r'\{\{\s*' + re.escape(col) + r'\s*\}\}', re.IGNORECASE),
                    re.compile(r'\{' + re.escape(col) + r'\}', re.IGNORECASE),
                    re.compile(r'\[' + re.escape(col) + r'\]', re.IGNORECASE),
                ]
                for pattern in patterns:
                    rendered_xml = pattern.sub(val, rendered_xml)

            clean_id, clean_name = get_document_clean_names(row, idx, list(df.columns))
            filename = f"DOC_{idx+1:04d}_{clean_id}{clean_name}.xml"

            zf.writestr(filename, rendered_xml.encode('utf-8'))
            total_generated += 1

    zip_buffer.seek(0)
    return zip_buffer, total_generated


def generate_batch_pdf_zip(template_bytes: bytes, df: pd.DataFrame) -> tuple[io.BytesIO, int]:
    reg_path = ROBOTO_FONTS['regular']
    zip_buffer = io.BytesIO()
    total_generated = 0


    normalized_template, _ = convert_pdf_all_to_roboto(template_bytes)

    with zipfile.ZipFile(zip_buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        for idx, row in df.iterrows():
            doc = pymupdf.open(stream=normalized_template, filetype="pdf")

            for page in doc:
                replacements_to_draw = []

                for col in df.columns:
                    val = str(row[col]) if row[col] is not None else ""
                    placeholder_variants = [
                        f"{{{{{col}}}}}",
                        f"{{{{ {col} }}}}",
                        f"{{{col}}}",
                        f"[{col}]"
                    ]

                    for p_var in placeholder_variants:
                        rects = page.search_for(p_var)
                        for r in rects:
                            page.add_redact_annot(r, fill=None)
                            replacements_to_draw.append((r, val))

                page.apply_redactions(images=0, graphics=0, text=0)

                for r, val in replacements_to_draw:
                    point = pymupdf.Point(r.x0, r.y1 - 2)
                    if os.path.exists(reg_path):
                        page.insert_text(point, val, fontname="Roboto-Regular", fontfile=reg_path, fontsize=10, color=(0.0, 0.17, 0.42))
                    else:
                        page.insert_text(point, val, fontname="helv", fontsize=10, color=(0.0, 0.17, 0.42))

            clean_id, clean_name = get_document_clean_names(row, idx, list(df.columns))
            filename = f"DOC_{idx+1:04d}_{clean_id}{clean_name}.pdf"

            try:
                doc.subset_fonts()
            except Exception:
                pass

            zf.writestr(filename, doc.tobytes(garbage=4, deflate=True))
            doc.close()
            total_generated += 1

    zip_buffer.seek(0)
    return zip_buffer, total_generated


@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    response = templates.TemplateResponse(request=request, name="index.html")
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.get("/favicon.ico", include_in_schema=False)
async def favicon():
    svg_icon = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'
        '<rect width="100" height="100" rx="24" fill="#2563EB"/>'
        '<text x="50" y="65" font-family="system-ui, sans-serif" font-weight="bold" '
        'font-size="52" fill="#FFFFFF" text-anchor="middle">F</text>'
        '</svg>'
    )
    return Response(content=svg_icon, media_type="image/svg+xml")


def process_single_file_content(file_bytes: bytes, filename: str, target_font: str = "roboto") -> tuple[bytes, str, int]:
    ext = filename.lower().split('.')[-1]
    if ext == 'xml':
        converted_bytes, replacements = convert_xml_to_font(file_bytes, target_font)
        media_type = "application/xml"
    elif ext == 'docx':
        converted_bytes, replacements = convert_docx_to_font(file_bytes, target_font)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif ext == 'pdf':
        converted_bytes, replacements = convert_pdf_to_font(file_bytes, target_font)
        media_type = "application/pdf"
    elif ext == 'rtf':
        converted_bytes, replacements = convert_rtf_to_font(file_bytes, target_font)
        media_type = "application/rtf"
    elif ext == 'doc':
        converted_bytes, replacements = convert_doc_to_font(file_bytes, target_font)
        media_type = "application/msword"
    else:
        raise HTTPException(
            status_code=400,
            detail=f"Tipe file '{filename}' tidak didukung. Format yang didukung: .xml, .docx, .doc, .pdf, .rtf, atau .zip."
        )
    return converted_bytes, media_type, replacements


@app.get("/api/fonts")
async def get_fonts_list():
    font_list = []
    for k, v in FONTS_REGISTRY.items():
        font_list.append({
            "id": v["id"],
            "name": v["name"],
            "family": v["family"],
            "prefix": v["prefix"],
            "badge": v["badge"],
            "description": v["description"]
        })
    return {"fonts": font_list, "default": "roboto"}


@app.post("/api/convert-font")
async def convert_font_endpoint(
    file: Optional[UploadFile] = File(None),
    files: Optional[list[UploadFile]] = File(None),
    target_font: Optional[str] = Form("roboto")
):
    upload_list: list[UploadFile] = []
    if files:
        upload_list.extend(files)
    if file:
        upload_list.append(file)

    if not upload_list:
        raise HTTPException(status_code=400, detail="Tidak ada file yang diunggah.")

    selected_font = target_font or "roboto"
    font_cfg = get_target_font(selected_font)
    prefix = font_cfg['prefix']

    if len(upload_list) == 1 and not (upload_list[0].filename or "").lower().endswith('.zip'):
        single_file = upload_list[0]
        filename = single_file.filename or "document"
        file_bytes = await single_file.read()
        if not file_bytes:
            raise HTTPException(status_code=400, detail="File kosong.")

        converted_bytes, media_type, replacements = process_single_file_content(file_bytes, filename, selected_font)
        out_name = f"{prefix}{filename}"

        return StreamingResponse(
            io.BytesIO(converted_bytes),
            media_type=media_type,
            headers={
                "Content-Disposition": f'attachment; filename="{out_name}"',
                "X-Replacements-Count": str(replacements),
                "X-Target-Font": font_cfg["name"],
                "Access-Control-Expose-Headers": "Content-Disposition, X-Replacements-Count, X-Target-Font"
            }
        )

    out_zip_buffer = io.BytesIO()
    total_converted = 0
    total_replacements = 0
    used_names = set()

    def get_unique_zip_entry(name: str) -> str:
        base = name
        counter = 1
        while name in used_names:
            parts = base.rsplit('.', 1)
            if len(parts) == 2:
                name = f"{parts[0]}_{counter}.{parts[1]}"
            else:
                name = f"{base}_{counter}"
            counter += 1
        used_names.add(name)
        return name

    with zipfile.ZipFile(out_zip_buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as zf:
        for ufile in upload_list:
            fname = ufile.filename or "file"
            fbytes = await ufile.read()
            if not fbytes:
                continue

            f_ext = fname.lower().split('.')[-1]
            if f_ext == 'zip':
                try:
                    with zipfile.ZipFile(io.BytesIO(fbytes)) as in_zf:
                        for item in in_zf.infolist():
                            if item.is_dir() or item.filename.startswith('__MACOSX') or item.filename.split('/')[-1].startswith('.'):
                                continue
                            item_ext = item.filename.lower().split('.')[-1]
                            if item_ext in ['xml', 'docx', 'doc', 'pdf', 'rtf']:
                                raw_item_bytes = in_zf.read(item)
                                c_bytes, _, reps = process_single_file_content(raw_item_bytes, item.filename, selected_font)
                                parts = item.filename.split('/')
                                parts[-1] = f"{prefix}{parts[-1]}"
                                zip_entry_name = "/".join(parts)
                                zf.writestr(zip_entry_name, c_bytes)
                                total_converted += 1
                                total_replacements += reps
                except Exception as e:
                    raise HTTPException(status_code=400, detail=f"Gagal memproses file ZIP '{fname}': {str(e)}")
            elif f_ext in ['xml', 'docx', 'doc', 'pdf', 'rtf']:
                c_bytes, _, reps = process_single_file_content(fbytes, fname, selected_font)
                entry_name = get_unique_zip_entry(f"{prefix}{fname}")
                zf.writestr(entry_name, c_bytes)
                total_converted += 1
                total_replacements += reps

    if total_converted == 0:
        raise HTTPException(status_code=400, detail="Tidak ditemukan file PDF, DOCX, DOC, XML, atau RTF yang valid untuk dikonversi.")

    out_zip_buffer.seek(0)

    if len(upload_list) == 1 and (upload_list[0].filename or "").lower().endswith('.zip'):
        raw_name = upload_list[0].filename
        base_zip = os.path.splitext(raw_name)[0]
        out_zip_name = f"{prefix}{base_zip}.zip"
    else:
        out_zip_name = f"{prefix}Converted_{total_converted}_Files.zip"

    return StreamingResponse(
        out_zip_buffer,
        media_type="application/zip",
        headers={
            "Content-Disposition": f'attachment; filename="{out_zip_name}"',
            "X-Converted-Count": str(total_converted),
            "X-Replacements-Count": str(total_replacements),
            "X-Target-Font": font_cfg["name"],
            "Access-Control-Expose-Headers": "Content-Disposition, X-Converted-Count, X-Replacements-Count, X-Target-Font"
        }
    )


@app.post("/api/batch-merge")
async def batch_merge_endpoint(
    template_file: UploadFile = File(...),
    data_file: UploadFile = File(...)
):
    t_filename = template_file.filename or ""
    t_ext = t_filename.lower().split('.')[-1]
    if t_ext not in ['xml', 'pdf']:
        raise HTTPException(status_code=400, detail="File Master Template harus berformat XML (.xml) atau PDF (.pdf).")

    d_filename = data_file.filename or ""
    d_ext = d_filename.lower().split('.')[-1]
    if d_ext not in ['csv', 'xlsx', 'xls']:
        raise HTTPException(status_code=400, detail="File Data Nasabah harus berformat CSV atau Excel (.xlsx/.xls).")

    template_bytes = await template_file.read()
    data_bytes = await data_file.read()

    if not template_bytes or not data_bytes:
        raise HTTPException(status_code=400, detail="Template atau Data Nasabah tidak boleh kosong.")

    df = load_dataset(data_bytes, d_filename)
    if df.empty:
        raise HTTPException(status_code=400, detail="File Data Nasabah tidak memiliki baris data (kosong).")

    if t_ext == 'xml':
        zip_buffer, count = generate_batch_xml_zip(template_bytes, df)
        out_zip_name = f"AstraLife_XML_Batch_Result_{count}_Docs.zip"
    else:
        zip_buffer, count = generate_batch_pdf_zip(template_bytes, df)
        out_zip_name = f"AstraLife_PDF_Batch_Result_{count}_Docs.zip"

    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={
            "Content-Disposition": f'attachment; filename="{out_zip_name}"',
            "X-Generated-Count": str(count),
            "Access-Control-Expose-Headers": "Content-Disposition, X-Generated-Count"
        }
    )


@app.post("/api/preview-data")
async def preview_data_endpoint(data_file: UploadFile = File(...)):
    filename = data_file.filename or ""
    file_bytes = await data_file.read()
    df = load_dataset(file_bytes, filename)

    columns = list(df.columns)
    preview_rows = df.head(5).to_dict(orient="records")
    total_rows = len(df)

    return {
        "columns": columns,
        "rows": preview_rows,
        "total_rows": total_rows
    }


@app.get("/api/download-sample/{sample_type}")
async def download_sample(sample_type: str):
    if sample_type == "xml":
        path = os.path.join(SAMPLES_DIR, "template.xml")
        return FileResponse(path, filename="template_sample.xml", media_type="application/xml")
    elif sample_type == "docx":
        path = os.path.join(SAMPLES_DIR, "template.docx")
        return FileResponse(path, filename="template_sample.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    elif sample_type == "pdf":
        path = os.path.join(SAMPLES_DIR, "template.pdf")
        return FileResponse(path, filename="template_sample.pdf", media_type="application/pdf")
    elif sample_type == "csv":
        path = os.path.join(SAMPLES_DIR, "data_nasabah.csv")
        return FileResponse(path, filename="data_nasabah_sample.csv", media_type="text/csv")
    else:
        raise HTTPException(status_code=404, detail="Tipe sample tidak ditemukan.")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
