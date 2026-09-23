import io
import os
import zipfile
import docx
import pymupdf
from fastapi.testclient import TestClient

from main import (
    app,
    convert_xml_arial_to_roboto,
    convert_xml_to_font,
    convert_docx_arial_to_roboto,
    convert_docx_to_font,
    convert_pdf_all_to_roboto,
    convert_pdf_to_font,
    load_dataset,
    generate_batch_xml_zip,
    generate_batch_pdf_zip,
    ROBOTO_FONTS,
    FONTS_REGISTRY,
    get_target_font
)

def test_backend_logic():
    with open("assets/samples/template.xml", "rb") as f:
        xml_bytes = f.read()
    converted_xml, count = convert_xml_arial_to_roboto(xml_bytes)
    assert count > 0
    assert b"Roboto" in converted_xml
    assert b"Arial" not in converted_xml
    print(f"XML Font Conversion OK: {count} replaced.")

    with open("assets/samples/template.docx", "rb") as f:
        docx_bytes = f.read()
    converted_docx, docx_count = convert_docx_arial_to_roboto(docx_bytes)
    assert docx_count > 0
    print(f"DOCX Font Conversion OK: {docx_count} replaced.")

    with open("assets/samples/template.pdf", "rb") as f:
        pdf_bytes = f.read()
    converted_pdf, pdf_count = convert_pdf_all_to_roboto(pdf_bytes)
    assert len(converted_pdf) > 0
    assert pdf_count > 0

    doc_chk = pymupdf.open(stream=converted_pdf, filetype="pdf")
    pdf_fonts = doc_chk[0].get_fonts()
    print("PDF Fonts in output:", [(f[0], f[3], f[4]) for f in pdf_fonts])

    roboto_present = any("Roboto" in f[3] for f in pdf_fonts)
    arial_present = any("Arial" in f[3] for f in pdf_fonts)

    assert roboto_present, f"Expected Roboto font, found: {pdf_fonts}"
    assert not arial_present, f"Arial font should not remain: {pdf_fonts}"
    doc_chk.close()
    print(f"PDF Font Conversion OK: {pdf_count} spans processed.")


    doc_times = pymupdf.open()
    p_t = doc_times.new_page()
    p_t.insert_text((50, 50), "Times Roman Text", fontname="tiro", fontsize=12)
    p_t.insert_text((50, 80), "Times Bold Text", fontname="tibo", fontsize=14)
    times_pdf_bytes = doc_times.tobytes()
    doc_times.close()

    converted_times_pdf, times_count = convert_pdf_all_to_roboto(times_pdf_bytes)
    chk_times = pymupdf.open(stream=converted_times_pdf, filetype="pdf")
    fonts_times = chk_times[0].get_fonts()
    for f in fonts_times:
        assert "Times" not in f[3], f"Times should not remain: {f[3]}"
        assert "Roboto" in f[3] or "Roboto" in f[4]
    chk_times.close()
    print("Times PDF Font Conversion OK: No Times font remains in output.")

    doc_bullet = pymupdf.open()
    p_b = doc_bullet.new_page()
    arial_path = "C:/Windows/Fonts/arial.ttf" if os.path.exists("C:/Windows/Fonts/arial.ttf") else ROBOTO_FONTS['regular']
    p_b.insert_font(fontname="Arial", fontfile=arial_path, set_simple=False)
    p_b.insert_text((50, 50), "\u25aa Pay special attention", fontname="Arial", fontsize=10)
    p_b.insert_text((50, 70), "\u25aa Check accuracy", fontname="Arial", fontsize=10)
    p_b.insert_text((50, 90), "1. Crossref.org", fontname="Arial", fontsize=10)
    bullet_pdf_bytes = doc_bullet.tobytes()
    doc_bullet.close()

    converted_bullet_pdf, bullet_count = convert_pdf_all_to_roboto(bullet_pdf_bytes)
    assert len(converted_bullet_pdf) > 0
    assert bullet_count >= 3
    chk_bullet = pymupdf.open(stream=converted_bullet_pdf, filetype="pdf")
    drawings = chk_bullet[0].get_drawings()
    assert len(drawings) >= 2, "Square bullets should be rendered as vector rectangles"
    for f in chk_bullet[0].get_fonts():
        assert "Roboto" in f[3] or "Roboto" in f[4]
        assert "Arial" not in f[3]
    chk_bullet.close()
    print("Bullet Symbol PDF Conversion OK: Square vector bullets rendered properly and Arial replaced by Roboto.")

    doc_dup = pymupdf.open()
    p_dup = doc_dup.new_page()
    p_dup.insert_text((50, 100), "Fitur Utama Asuransi", fontname="helv", fontsize=8.8)
    p_dup.draw_rect(pymupdf.Rect(40, 85, 200, 115), fill=(0, 0.3, 0.6))
    p_dup.insert_text((49, 98), "Fitur Utama Asuransi", fontname="helv", fontsize=7.0)
    dup_pdf_bytes = doc_dup.tobytes()
    doc_dup.close()

    converted_dup_pdf, dup_count = convert_pdf_all_to_roboto(dup_pdf_bytes)
    chk_dup = pymupdf.open(stream=converted_dup_pdf, filetype="pdf")
    fitur_occurrences = []
    for b in chk_dup[0].get_text("dict")["blocks"]:
        if b.get("type") == 0:
            for l in b.get("lines", []):
                for s in l.get("spans", []):
                    if "Fitur Utama Asuransi" in s.get("text", ""):
                        fitur_occurrences.append(s)
    assert len(fitur_occurrences) == 1, f"Expected 1 deduplicated span, found {len(fitur_occurrences)}"
    chk_dup.close()
    print("Overlapping Span Deduplication OK: Ghost/covered text suppressed successfully.")

    doc_space = pymupdf.open()
    p_sp = doc_space.new_page()
    p_sp.insert_text((50, 50), "Anda", fontname="hebo", fontsize=7.0)
    p_sp.insert_text((65, 50), " adalah Tertanggung", fontname="helv", fontsize=7.0)
    space_pdf_bytes = doc_space.tobytes()
    doc_space.close()

    converted_sp_pdf, _ = convert_pdf_all_to_roboto(space_pdf_bytes)
    chk_sp = pymupdf.open(stream=converted_sp_pdf, filetype="pdf")
    line_spans = chk_sp[0].get_text("dict")["blocks"][0]["lines"][0]["spans"]
    assert len(line_spans) == 2
    assert "Roboto" in line_spans[0]["font"] and "Bold" in line_spans[0]["font"]
    assert line_spans[0]["text"] == "Anda"
    assert "Roboto" in line_spans[1]["font"]
    assert "Andaadalah" not in chk_sp[0].get_text("text")
    assert "Anda adalah" in chk_sp[0].get_text("text")
    chk_sp.close()
    print("Inline Spacing and Space Preservation OK: No merged words, Bold and Regular styles preserved.")

    doc_fit = pymupdf.open()
    p_fit = doc_fit.new_page()
    p_fit.insert_text((50, 100), "dengan lengkap dan benar; dan dapat diperpanjang", fontname="helv", fontsize=6.0)
    fit_pdf_bytes = doc_fit.tobytes()
    doc_fit.close()

    converted_fit_pdf, _ = convert_pdf_all_to_roboto(fit_pdf_bytes)
    chk_fit = pymupdf.open(stream=converted_fit_pdf, filetype="pdf")
    fit_span = chk_fit[0].get_text("dict")["blocks"][0]["lines"][0]["spans"][0]
    assert "Roboto" in fit_span["font"]
    assert fit_span["size"] <= 6.0
    chk_fit.close()
    print("Auto-Fit Font Scaling OK: Font size scaled to fit original line boundaries.")

    doc_just = pymupdf.open()
    p_just = doc_just.new_page()
    p_just.insert_text((50, 100), "Sertifikat Asuransi ini menjelaskan perlindungan", fontname="helv", fontsize=8.0)
    p_just.insert_text((50, 115), "singkat serta hak kewajiban sehubungan perjanjian", fontname="helv", fontsize=8.0)
    p_just.insert_text((50, 130), "disembunyikan oleh Peserta (Tertanggung) dan", fontname="helv", fontsize=8.0)
    just_pdf_bytes = doc_just.tobytes()
    doc_just.close()
    converted_just_pdf, _ = convert_pdf_all_to_roboto(just_pdf_bytes)
    chk_just = pymupdf.open(stream=converted_just_pdf, filetype="pdf")
    assert len(chk_just) == 1
    d_just = chk_just[0].get_text("rawdict")
    for b in d_just["blocks"]:
        for l in b.get("lines", []):
            all_c = [c for s in l["spans"] for c in s["chars"]]
            for i in range(len(all_c) - 1):
                assert all_c[i+1]["bbox"][0] - all_c[i]["bbox"][2] <= 5.0
    chk_just.close()
    print("Justified Text Alignment OK: Word spacing preserved on justified paragraphs.")

    doc_num = pymupdf.open()
    p_num = doc_num.new_page()
    p_num.insert_text((40, 100), "1.", fontname="helv", fontsize=8.0)
    p_num.insert_text((51, 100), "Ketentuan umum mengenai polis dan sertifikat", fontname="helv", fontsize=8.0)
    p_num.insert_text((54, 115), "sebagaimana diatur dalam ketentuan yang berlaku", fontname="helv", fontsize=8.0)
    num_pdf_bytes = doc_num.tobytes()
    doc_num.close()
    converted_num_pdf, _ = convert_pdf_all_to_roboto(num_pdf_bytes)
    chk_num = pymupdf.open(stream=converted_num_pdf, filetype="pdf")
    d_num = chk_num[0].get_text("dict")
    found_aligned = False
    for b in d_num["blocks"]:
        if b.get("type") == 0:
            for l in b["lines"]:
                txt = "".join(s["text"] for s in l["spans"]).strip()
                if "Ketentuan" in txt:
                    for s in l["spans"]:
                        if "Ketentuan" in s["text"]:
                            assert abs(s["origin"][0] - 54.0) < 1.0
                            found_aligned = True
    assert found_aligned
    chk_num.close()
    print("Numbered List Indentation Alignment OK: First line body aligns with continuation indent.")

    csv_path = "assets/samples/data_nasabah.csv"
    with open(csv_path, "rb") as f:
        csv_bytes = f.read()
    df = load_dataset(csv_bytes, "data_sample.csv")
    assert len(df) == 6

    zip_buffer, total_gen = generate_batch_xml_zip(converted_xml, df)
    assert total_gen == 6
    with zipfile.ZipFile(zip_buffer, 'r') as zf:
        namelist = zf.namelist()
        assert len(namelist) == 6
    print("XML Batch ZIP OK.")

    pdf_zip_buffer, pdf_total_gen = generate_batch_pdf_zip(pdf_bytes, df)
    assert pdf_total_gen == 6
    with zipfile.ZipFile(pdf_zip_buffer, 'r') as zf:
        namelist = zf.namelist()
        assert len(namelist) == 6
    print("PDF Batch ZIP OK.")

    client = TestClient(app)

    res = client.get("/")
    assert res.status_code == 200
    assert "Converting Font" in res.text
    print("GET / OK")

    res = client.post("/api/convert-font", files={"file": ("test.xml", xml_bytes, "application/xml")})
    assert res.status_code == 200
    res = client.post("/api/convert-font", files={"file": ("test.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res.status_code == 200
    res = client.post("/api/convert-font", files={"file": ("test.pdf", pdf_bytes, "application/pdf")})
    assert res.status_code == 200


    zip_in = io.BytesIO()
    with zipfile.ZipFile(zip_in, "w") as zf:
        zf.writestr("sub/doc1.pdf", pdf_bytes)
        zf.writestr("doc2.xml", xml_bytes)
    res_zip = client.post("/api/convert-font", files={"file": ("archive.zip", zip_in.getvalue(), "application/zip")})
    assert res_zip.status_code == 200
    assert res_zip.headers["content-type"] == "application/zip"
    assert int(res_zip.headers.get("x-converted-count", 0)) == 2
    with zipfile.ZipFile(io.BytesIO(res_zip.content)) as out_zf:
        assert "sub/Roboto_doc1.pdf" in out_zf.namelist()
        assert "Roboto_doc2.xml" in out_zf.namelist()


    multi_payload = [
        ("files", ("f1.pdf", pdf_bytes, "application/pdf")),
        ("files", ("f2.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ]
    res_multi = client.post("/api/convert-font", files=multi_payload)
    assert res_multi.status_code == 200
    assert res_multi.headers["content-type"] == "application/zip"
    assert int(res_multi.headers.get("x-converted-count", 0)) == 2
    print("POST /api/convert-font (Single, ZIP, and Multi-file) OK")

    res = client.post("/api/preview-data", files={"data_file": ("data.csv", csv_bytes, "text/csv")})
    assert res.status_code == 200
    assert res.json()["total_rows"] == 6
    print("POST /api/preview-data OK")

    res = client.post("/api/batch-merge", files={"template_file": ("template.xml", xml_bytes, "application/xml"), "data_file": ("data.csv", csv_bytes, "text/csv")})
    assert res.status_code == 200
    assert int(res.headers.get("x-generated-count")) == 6
    print("POST /api/batch-merge (XML) OK")

    res = client.post("/api/batch-merge", files={"template_file": ("template.pdf", pdf_bytes, "application/pdf"), "data_file": ("data.csv", csv_bytes, "text/csv")})
    assert res.status_code == 200
    assert int(res.headers.get("x-generated-count")) == 6
    print("POST /api/batch-merge (PDF) OK")

    for sample in ["pdf", "xml", "docx", "csv"]:
        res = client.get(f"/api/download-sample/{sample}")
        assert res.status_code == 200
        print(f"GET /api/download-sample/{sample} OK")

    res_fonts = client.get("/api/fonts")
    assert res_fonts.status_code == 200
    font_data = res_fonts.json()
    assert font_data["default"] == "roboto"
    available_ids = [f["id"] for f in font_data["fonts"]]
    for expected_id in ["roboto", "opensans", "montserrat", "arial", "times", "calibri", "segoeui"]:
        assert expected_id in available_ids, f"Font {expected_id} should be available"
    print("GET /api/fonts OK: All 7 target fonts registered and exposed.")

    converted_opensans_xml, count_os = convert_xml_to_font(xml_bytes, "opensans")
    assert count_os > 0
    assert b"Open Sans" in converted_opensans_xml
    assert b"Arial" not in converted_opensans_xml
    print("XML to Open Sans OK.")

    converted_mont_xml, count_mont = convert_xml_to_font(xml_bytes, "montserrat")
    assert count_mont > 0
    assert b"Montserrat" in converted_mont_xml
    assert b"Arial" not in converted_mont_xml
    print("XML to Montserrat OK.")

    custom_xml = b'<?xml version="1.0"?><doc><item font-family="Calibri">Calibri Text</item><item font-family="Times New Roman">Times Text</item></doc>'
    converted_custom_xml, count_c = convert_xml_to_font(custom_xml, "roboto")
    assert count_c >= 2
    assert b"Roboto" in converted_custom_xml
    assert b"Calibri" not in converted_custom_xml
    assert b"Times New Roman" not in converted_custom_xml
    print("Universal Source Font XML OK: Calibri and Times New Roman converted to Roboto.")

    custom_docx = docx.Document()
    p1 = custom_docx.add_paragraph()
    r1 = p1.add_run("Calibri Run")
    r1.font.name = "Calibri"
    p2 = custom_docx.add_paragraph()
    r2 = p2.add_run("Times Run")
    r2.font.name = "Times New Roman"
    docx_buf = io.BytesIO()
    custom_docx.save(docx_buf)
    converted_multi_docx, docx_m_count = convert_docx_to_font(docx_buf.getvalue(), "montserrat")
    assert docx_m_count >= 2
    doc_m_chk = docx.Document(io.BytesIO(converted_multi_docx))
    for p in doc_m_chk.paragraphs:
        for r in p.runs:
            assert r.font.name == "Montserrat"
    print("Universal Source Font DOCX OK: Calibri and Times New Roman converted to Montserrat.")

    converted_os_pdf, os_p_count = convert_pdf_to_font(pdf_bytes, "opensans")
    assert os_p_count > 0
    doc_os = pymupdf.open(stream=converted_os_pdf, filetype="pdf")
    os_fonts = doc_os[0].get_fonts()
    assert any("OpenSans" in f[3] or "OpenSans" in f[4] for f in os_fonts)
    doc_os.close()
    print("PDF to Open Sans OK: OpenSans fonts applied.")

    converted_mont_pdf, mont_p_count = convert_pdf_to_font(pdf_bytes, "montserrat")
    assert mont_p_count > 0
    doc_mont = pymupdf.open(stream=converted_mont_pdf, filetype="pdf")
    mont_fonts = doc_mont[0].get_fonts()
    assert any("Montserrat" in f[3] or "Montserrat" in f[4] for f in mont_fonts)
    doc_mont.close()
    print("PDF to Montserrat OK: Montserrat fonts applied.")

    res_target_api = client.post(
        "/api/convert-font",
        files={"file": ("report.pdf", pdf_bytes, "application/pdf")},
        data={"target_font": "montserrat"}
    )
    assert res_target_api.status_code == 200
    assert 'filename="Montserrat_report.pdf"' in res_target_api.headers.get("content-disposition", "")
    assert res_target_api.headers.get("x-target-font") == "Montserrat"
    print("POST /api/convert-font with target_font=montserrat OK.")

    neo_xml = b'<?xml version="1.0"?><doc><item font-family="Neo Sans Pro">Neo Sans Text</item><item font-family="NeoSans">NeoSans Text</item><item font-family="NeoSansPro">NeoSansPro Text</item></doc>'
    converted_neo_xml, count_neo = convert_xml_to_font(neo_xml, 'roboto')
    assert count_neo >= 3, f"Expected >= 3 replacements, got {count_neo}"
    assert b"Roboto" in converted_neo_xml
    assert b"Neo Sans" not in converted_neo_xml
    assert b"NeoSans" not in converted_neo_xml
    print("Neo Sans XML -> Roboto OK: All Neo Sans variants replaced.")

    neo_docx = docx.Document()
    p_neo1 = neo_docx.add_paragraph()
    r_neo1 = p_neo1.add_run("Neo Sans Pro Run")
    r_neo1.font.name = "Neo Sans Pro"
    p_neo2 = neo_docx.add_paragraph()
    r_neo2 = p_neo2.add_run("NeoSans Run")
    r_neo2.font.name = "NeoSans"
    neo_docx_buf = io.BytesIO()
    neo_docx.save(neo_docx_buf)
    converted_neo_docx, neo_docx_count = convert_docx_to_font(neo_docx_buf.getvalue(), "roboto")
    assert neo_docx_count >= 2, f"Expected >= 2 DOCX replacements, got {neo_docx_count}"
    doc_neo_chk = docx.Document(io.BytesIO(converted_neo_docx))
    for p in doc_neo_chk.paragraphs:
        for r in p.runs:
            assert r.font.name == "Roboto", f"Expected Roboto, got {r.font.name}"
    print("Neo Sans DOCX -> Roboto OK: Neo Sans Pro and NeoSans runs converted.")

    auto_cfg = get_target_font("auto")
    assert auto_cfg["family"] == "Roboto", f"Auto mode must map to Roboto, got {auto_cfg['family']}"
    auto_cfg2 = get_target_font("")
    assert auto_cfg2["family"] == "Roboto"
    auto_cfg3 = get_target_font(None)
    assert auto_cfg3["family"] == "Roboto"
    print("Auto Mode Mapping OK: 'auto', '', None all resolve to Roboto.")

    print("ALL TESTS PASSED!")

if __name__ == "__main__":
    test_backend_logic()
